"""Task template management.

A project's template is the ordered list of activities each node must go
through. Uploading it materialises one Task row per (node × template activity) —
2,850 rows for a 57-node project.

The legacy uploader (`routes/task.py`) was unauthenticated, deleted every task
in the project on each upload, and left the corresponding TaskExecution rows
behind pointing at task ids that no longer existed. That is the mechanism that
produced the 193 orphan audit rows found during migration (audit B5, B6, M19).

Here the template is reconciled rather than replaced: activities are matched by
template number, existing rows are updated in place, and any activity carrying
recorded dates cannot be removed without an explicit force.
"""
from __future__ import annotations

from io import BytesIO

from fastapi import (
    APIRouter,
    Depends,
    File,
    HTTPException,
    Query,
    Response,
    UploadFile,
    status,
)
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.api.common import read_upload, require_project, xlsx_response
from app.db import get_db
from app.services.baseline import maybe_autobaseline
from app.models import AuditLog, Scope, Task, TaskExecution
from app.security import CurrentUser, get_current_user, require_admin, require_csrf

router = APIRouter()


def _autobaseline_note(db, project_id: int, actor: str) -> dict | None:
    """Plan the project if this save was the one that made it plannable.

    Returned to the caller so the UI can say a plan was generated, rather than
    the dates simply appearing with no explanation.
    """
    result = maybe_autobaseline(db, project_id, actor=actor)
    if result is None:
        return None
    return {
        "generated": True,
        "baseline_version": result.baseline_version,
        "tasks_planned": result.tasks_planned,
        "plan_start": result.plan_start.isoformat() if result.plan_start else None,
        "plan_finish": result.plan_finish.isoformat() if result.plan_finish else None,
    }

REQUIRED_COLUMNS = ["Task_Number", "Task_Name", "Duration_Days", "Predecessor_Task_Number"]
OPTIONAL_COLUMNS = ["Is_Prerequisite", "Owner_Role"]


class TemplateRow(BaseModel):
    template_task_number: int
    name: str
    duration_days: int
    predecessor_template_number: int | None
    is_prerequisite: bool
    owner_role: str | None
    node_count: int
    recorded_dates: int


@router.get("/projects/{project_id}/template", response_model=list[TemplateRow])
def list_template(
    project_id: int,
    db: Session = Depends(get_db),
    _: CurrentUser = Depends(get_current_user),
):
    require_project(db, project_id)
    rows = db.execute(
        select(
            Task.template_task_number,
            func.min(Task.name),
            func.min(Task.duration_days),
            func.min(Task.predecessor_template_number),
            func.bool_or(Task.is_prerequisite),
            func.min(Task.owner_role),
            func.count(func.distinct(Task.scope_id)),
            func.count(TaskExecution.id).filter(TaskExecution.actual_start.is_not(None)),
        )
        .join(TaskExecution, TaskExecution.task_id == Task.id, isouter=True)
        .where(Task.project_id == project_id)
        .group_by(Task.template_task_number)
        .order_by(Task.template_task_number)
    ).all()
    return [
        TemplateRow(
            template_task_number=number,
            name=name,
            duration_days=duration,
            predecessor_template_number=predecessor,
            is_prerequisite=bool(prerequisite),
            owner_role=owner,
            node_count=nodes,
            recorded_dates=recorded,
        )
        for number, name, duration, predecessor, prerequisite, owner, nodes, recorded in rows
    ]


@router.get("/projects/{project_id}/template/export")
def export_template(
    project_id: int,
    db: Session = Depends(get_db),
    _: CurrentUser = Depends(get_current_user),
):
    """Download the current template, or a starter sheet if none is loaded."""
    require_project(db, project_id)
    rows = list_template(project_id, db, _)  # type: ignore[arg-type]

    wb = Workbook()
    ws = wb.active
    ws.title = "Template"
    ws.append(REQUIRED_COLUMNS + OPTIONAL_COLUMNS)
    if rows:
        for row in rows:
            ws.append([
                row.template_task_number, row.name, row.duration_days,
                row.predecessor_template_number, "Y" if row.is_prerequisite else "",
                row.owner_role or "",
            ])
    else:
        ws.append([1, "Scope defined and available", 0, None, "Y", "Planning"])
        ws.append([2, "Project Initiation", 1, 1, "", "Planning"])
    for index, width in enumerate([14, 42, 15, 24, 16, 16], start=1):
        ws.column_dimensions[chr(64 + index)].width = width
    ws.freeze_panes = "A2"

    guide = wb.create_sheet("How to use")
    for line in [
        "Task_Number  the activity's position in the template, 1..N. This is the",
        "             identifier everything else binds to — reporting, the",
        "             execution workbook and scheduling constraints. Do not",
        "             renumber an activity that already has recorded dates.",
        "Task_Name    what appears on screen and in exports.",
        "Duration_Days  working days. 0 marks a gate that consumes no time.",
        "Predecessor_Task_Number  another Task_Number in this sheet, or blank.",
        "Is_Prerequisite  Y for an activity satisfied at kickoff.",
        "Owner_Role   optional, free text.",
        "",
        "Uploading reconciles: activities are matched by Task_Number, existing",
        "rows are updated in place, and an activity carrying recorded dates",
        "cannot be removed without an explicit force.",
    ]:
        guide.append([line])
    guide.column_dimensions["A"].width = 78

    buffer = BytesIO()
    wb.save(buffer)
    return xlsx_response(buffer.getvalue(), f"IPTT_template_{project_id}.xlsx")


@router.post("/projects/{project_id}/template/import")
async def import_template(
    project_id: int,
    file: UploadFile = File(...),
    dry_run: bool = Query(True),
    force_remove: bool = Query(False, description="Remove activities that carry recorded dates"),
    db: Session = Depends(get_db),
    user: CurrentUser = Depends(require_admin),
    __: None = Depends(require_csrf),
):
    require_project(db, project_id)
    content = await read_upload(file)

    try:
        wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception:  # noqa: BLE001
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Not a readable .xlsx workbook")

    ws = wb["Template"] if "Template" in wb.sheetnames else wb[wb.sheetnames[0]]
    grid = list(ws.iter_rows(values_only=True))
    if not grid:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "The sheet is empty")

    header = [("" if v is None else str(v).strip()) for v in grid[0]]
    index_of = {h.lower(): i for i, h in enumerate(header) if h}
    missing = [c for c in REQUIRED_COLUMNS if c.lower() not in index_of]
    if missing:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            f"Missing column(s): {', '.join(missing)}",
        )

    errors: list[str] = []
    parsed: dict[int, dict] = {}

    for row_number, row in enumerate(grid[1:], start=2):
        def cell(name: str):
            i = index_of.get(name.lower())
            return row[i] if i is not None and i < len(row) else None

        raw_number = cell("Task_Number")
        if raw_number is None or str(raw_number).strip() == "":
            continue
        try:
            number = int(raw_number)
        except (TypeError, ValueError):
            errors.append(f"Row {row_number}: Task_Number '{raw_number}' is not a whole number")
            continue
        if not 1 <= number <= 999:
            errors.append(f"Row {row_number}: Task_Number {number} is outside 1..999")
            continue
        if number in parsed:
            errors.append(f"Row {row_number}: Task_Number {number} appears more than once")
            continue

        name = str(cell("Task_Name") or "").strip()
        if not name:
            errors.append(f"Row {row_number}: Task_Name is required")
            continue

        try:
            duration = int(cell("Duration_Days") or 0)
        except (TypeError, ValueError):
            errors.append(f"Row {row_number}: Duration_Days must be a whole number")
            continue
        if duration < 0:
            errors.append(f"Row {row_number}: Duration_Days cannot be negative")
            continue

        raw_predecessor = cell("Predecessor_Task_Number")
        predecessor: int | None = None
        if raw_predecessor not in (None, "", 0, "0"):
            try:
                predecessor = int(raw_predecessor)
            except (TypeError, ValueError):
                errors.append(
                    f"Row {row_number}: Predecessor_Task_Number '{raw_predecessor}' "
                    "is not a whole number"
                )
                continue

        prerequisite = str(cell("Is_Prerequisite") or "").strip().lower() in {"y", "yes", "true", "1"}
        owner = str(cell("Owner_Role") or "").strip() or None

        parsed[number] = {
            "template_task_number": number,
            "name": name,
            "duration_days": duration,
            "predecessor_template_number": predecessor,
            "is_prerequisite": prerequisite,
            "owner_role": owner,
        }

    if not parsed and not errors:
        errors.append("No activities found in the sheet")

    # Dependencies must resolve inside the template, and must not cycle.
    for number, values in parsed.items():
        predecessor = values["predecessor_template_number"]
        if predecessor is None:
            continue
        if predecessor not in parsed:
            errors.append(
                f"Activity {number} lists predecessor {predecessor}, which is not "
                "in the sheet"
            )
        elif predecessor == number:
            errors.append(f"Activity {number} is its own predecessor")

    if not errors:
        for number in parsed:
            seen: set[int] = set()
            cursor: int | None = number
            while cursor is not None:
                if cursor in seen:
                    errors.append(
                        f"Activities {sorted(seen)} form a dependency cycle"
                    )
                    break
                seen.add(cursor)
                cursor = parsed[cursor]["predecessor_template_number"]
            if errors:
                break

    # --- compare with what exists ------------------------------------------
    current = db.execute(
        select(
            Task.template_task_number,
            func.min(Task.name),
            func.count(TaskExecution.id).filter(TaskExecution.actual_start.is_not(None)),
        )
        .join(TaskExecution, TaskExecution.task_id == Task.id, isouter=True)
        .where(Task.project_id == project_id)
        .group_by(Task.template_task_number)
    ).all()
    existing = {number: (name, recorded) for number, name, recorded in current}

    removals = sorted(set(existing) - set(parsed))
    blocked = [
        f"{n} '{existing[n][0]}' ({existing[n][1]} recorded dates)"
        for n in removals
        if existing[n][1]
    ]
    if blocked and not force_remove:
        errors.append(
            "These activities are absent from the sheet but carry recorded dates: "
            + "; ".join(blocked[:10])
            + ". Re-send with force_remove=true to delete them and their history."
        )

    scope_count = db.scalar(
        select(func.count()).select_from(Scope).where(Scope.project_id == project_id)
    ) or 0

    summary = {
        "dry_run": dry_run,
        "ok": not errors,
        "activities_in_sheet": len(parsed),
        "nodes": scope_count,
        "to_create": sorted(set(parsed) - set(existing)),
        "to_update": sorted(set(parsed) & set(existing)),
        "to_remove": removals,
        "task_rows_affected": len(parsed) * scope_count,
        "errors": errors,
    }

    if errors:
        db.rollback()
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, summary)
    if dry_run:
        return summary
    if scope_count == 0:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            "Load the project scope before the task template — tasks are created "
            "per node.",
        )

    scopes = db.scalars(select(Scope).where(Scope.project_id == project_id)).all()
    tasks = db.scalars(select(Task).where(Task.project_id == project_id)).all()
    by_key = {(t.scope_id, t.template_task_number): t for t in tasks}

    for scope in scopes:
        for number, values in parsed.items():
            task = by_key.get((scope.id, number))
            if task is None:
                db.add(Task(project_id=project_id, scope_id=scope.id, **values))
            else:
                for key, value in values.items():
                    setattr(task, key, value)

    for number in removals:
        for task in [t for t in tasks if t.template_task_number == number]:
            db.delete(task)  # executions cascade

    db.add(
        AuditLog(
            actor_user_id=user.id, actor_username=user.username, actor_role=user.role,
            action="TEMPLATE_IMPORT", source="excel-import", project_id=project_id,
            field="task_template",
            new_value=(
                f"{len(parsed)} activities x {scope_count} nodes; "
                f"removed {len(removals)}"
            ),
        )
    )
    db.flush()
    summary["autobaseline"] = _autobaseline_note(db, project_id, user.username)
    return summary


# ---------------------------------------------------------------------------
# Row-level template editing
# ---------------------------------------------------------------------------
#
# A "template row" is not one database row. The template is materialised per
# node, so activity 32 of a 57-node project is 57 Task rows sharing
# template_task_number=32. Every operation here therefore fans out across the
# project's scopes, which is also why a template edit is a project-level
# action rather than something done to a single task.


class TemplateRowWrite(BaseModel):
    template_task_number: int = Field(ge=1, le=999)
    name: str = Field(min_length=1, max_length=200)
    duration_days: int = Field(ge=0, le=3650)
    predecessor_template_number: int | None = None
    is_prerequisite: bool = False
    owner_role: str | None = Field(default=None, max_length=100)


class TemplateRowUpdate(BaseModel):
    name: str | None = Field(default=None, min_length=1, max_length=200)
    duration_days: int | None = Field(default=None, ge=0, le=3650)
    predecessor_template_number: int | None = None
    is_prerequisite: bool | None = None
    owner_role: str | None = Field(default=None, max_length=100)


def _template_numbers(db: Session, project_id: int) -> dict[int, int | None]:
    """Every activity number in the template, mapped to its predecessor."""
    return {
        number: predecessor
        for number, predecessor in db.execute(
            select(Task.template_task_number, func.min(Task.predecessor_template_number))
            .where(Task.project_id == project_id)
            .group_by(Task.template_task_number)
        ).all()
    }


def _assert_dependencies_resolve(graph: dict[int, int | None]) -> None:
    """Same rules the sheet importer enforces, applied to a single edit.

    Without this, editing one row in the UI could introduce a dangling
    predecessor or a cycle that the importer would have rejected - and the
    planner would then either skip the activity or loop.
    """
    for number, predecessor in graph.items():
        if predecessor is None:
            continue
        if predecessor == number:
            raise HTTPException(
                status.HTTP_400_BAD_REQUEST,
                f"Activity {number} cannot be its own predecessor",
            )
        if predecessor not in graph:
            raise HTTPException(
                status.HTTP_400_BAD_REQUEST,
                f"Activity {number} lists predecessor {predecessor}, which is not "
                "in the template",
            )

    for number in graph:
        seen: set[int] = set()
        cursor: int | None = number
        while cursor is not None:
            if cursor in seen:
                raise HTTPException(
                    status.HTTP_400_BAD_REQUEST,
                    f"Activities {sorted(seen)} form a dependency cycle",
                )
            seen.add(cursor)
            cursor = graph[cursor]


@router.post("/projects/{project_id}/template", response_model=TemplateRow, status_code=201)
def add_template_row(
    project_id: int,
    payload: TemplateRowWrite,
    db: Session = Depends(get_db),
    user: CurrentUser = Depends(require_admin),
    _: None = Depends(require_csrf),
):
    require_project(db, project_id)

    graph = _template_numbers(db, project_id)
    if payload.template_task_number in graph:
        raise HTTPException(
            status.HTTP_409_CONFLICT,
            f"Activity {payload.template_task_number} already exists in this template",
        )
    graph[payload.template_task_number] = payload.predecessor_template_number
    _assert_dependencies_resolve(graph)

    scope_ids = list(
        db.scalars(select(Scope.id).where(Scope.project_id == project_id)).all()
    )
    if not scope_ids:
        raise HTTPException(
            status.HTTP_409_CONFLICT,
            "Add at least one node to the scope before building the task template",
        )

    for scope_id in scope_ids:
        db.add(
            Task(
                project_id=project_id,
                scope_id=scope_id,
                template_task_number=payload.template_task_number,
                name=payload.name.strip(),
                duration_days=payload.duration_days,
                predecessor_template_number=payload.predecessor_template_number,
                is_prerequisite=payload.is_prerequisite,
                owner_role=payload.owner_role,
            )
        )
    db.add(
        AuditLog(
            actor_user_id=user.id, actor_username=user.username, actor_role=user.role,
            action="TEMPLATE_ROW_CREATE", source="api", project_id=project_id,
            field="template_task_number", old_value=None,
            new_value=str(payload.template_task_number),
            task_name=payload.name.strip(),
        )
    )
    db.flush()
    _autobaseline_note(db, project_id, user.username)
    db.flush()
    return TemplateRow(
        template_task_number=payload.template_task_number,
        name=payload.name.strip(),
        duration_days=payload.duration_days,
        predecessor_template_number=payload.predecessor_template_number,
        is_prerequisite=payload.is_prerequisite,
        owner_role=payload.owner_role,
        node_count=len(scope_ids),
        recorded_dates=0,
    )


@router.patch("/projects/{project_id}/template/{number}", response_model=TemplateRow)
def update_template_row(
    project_id: int,
    number: int,
    payload: TemplateRowUpdate,
    db: Session = Depends(get_db),
    user: CurrentUser = Depends(require_admin),
    _: None = Depends(require_csrf),
):
    require_project(db, project_id)
    rows = list(
        db.scalars(
            select(Task).where(
                Task.project_id == project_id, Task.template_task_number == number
            )
        ).all()
    )
    if not rows:
        raise HTTPException(
            status.HTTP_404_NOT_FOUND, f"Activity {number} is not in this template"
        )

    changes = payload.model_dump(exclude_unset=True)
    if "predecessor_template_number" in changes:
        graph = _template_numbers(db, project_id)
        graph[number] = changes["predecessor_template_number"]
        _assert_dependencies_resolve(graph)
    if "name" in changes and changes["name"] is not None:
        changes["name"] = changes["name"].strip()

    for field, new in changes.items():
        old = getattr(rows[0], field)
        if old == new:
            continue
        for row in rows:           # the same edit across every node
            setattr(row, field, new)
        db.add(
            AuditLog(
                actor_user_id=user.id, actor_username=user.username, actor_role=user.role,
                action="TEMPLATE_ROW_UPDATE", source="api", project_id=project_id,
                field=field,
                old_value=None if old is None else str(old),
                new_value=None if new is None else str(new),
                task_name=rows[0].name,
            )
        )

    db.flush()
    recorded = db.scalar(
        select(func.count())
        .select_from(TaskExecution)
        .join(Task, Task.id == TaskExecution.task_id)
        .where(
            Task.project_id == project_id,
            Task.template_task_number == number,
            (TaskExecution.actual_start.is_not(None))
            | (TaskExecution.actual_finish.is_not(None)),
        )
    )
    head = rows[0]
    return TemplateRow(
        template_task_number=number,
        name=head.name,
        duration_days=head.duration_days,
        predecessor_template_number=head.predecessor_template_number,
        is_prerequisite=head.is_prerequisite,
        owner_role=head.owner_role,
        node_count=len(rows),
        recorded_dates=recorded or 0,
    )


@router.delete(
    "/projects/{project_id}/template/{number}",
    status_code=204,
    response_class=Response,
)
def delete_template_row(
    project_id: int,
    number: int,
    force: bool = Query(False),
    db: Session = Depends(get_db),
    user: CurrentUser = Depends(require_admin),
    _: None = Depends(require_csrf),
):
    """Removes the activity from every node.

    Refused while any node has recorded dates against it, unless forced - those
    are field observations, and the cascade to TaskExecution would take them
    with it. Also refused while another activity depends on it, which would
    otherwise leave a dangling predecessor the planner silently ignores.
    """
    require_project(db, project_id)
    rows = list(
        db.scalars(
            select(Task).where(
                Task.project_id == project_id, Task.template_task_number == number
            )
        ).all()
    )
    if not rows:
        raise HTTPException(
            status.HTTP_404_NOT_FOUND, f"Activity {number} is not in this template"
        )

    dependents = [
        n for n, predecessor in _template_numbers(db, project_id).items()
        if predecessor == number and n != number
    ]
    if dependents:
        raise HTTPException(
            status.HTTP_409_CONFLICT,
            f"Activities {sorted(dependents)} list {number} as their predecessor. "
            "Repoint them first.",
        )

    if not force:
        recorded = db.scalar(
            select(func.count())
            .select_from(TaskExecution)
            .join(Task, Task.id == TaskExecution.task_id)
            .where(
                Task.project_id == project_id,
                Task.template_task_number == number,
                (TaskExecution.actual_start.is_not(None))
                | (TaskExecution.actual_finish.is_not(None)),
            )
        )
        if recorded:
            raise HTTPException(
                status.HTTP_409_CONFLICT,
                f"{recorded} node(s) have recorded dates against activity {number}. "
                "Removing it destroys that field data. Pass force=true to proceed.",
            )

    db.add(
        AuditLog(
            actor_user_id=user.id, actor_username=user.username, actor_role=user.role,
            action="TEMPLATE_ROW_DELETE", source="api", project_id=project_id,
            field="template_task_number", old_value=str(number), new_value=None,
            task_name=rows[0].name,
        )
    )
    for row in rows:
        db.delete(row)
    return Response(status_code=204)
