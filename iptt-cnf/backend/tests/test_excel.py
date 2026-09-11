"""Excel round trip.

The centrepiece is `test_reordered_columns_still_import_correctly`. The legacy
importer matched activity columns by *position* relative to the "Facility"
column, so moving or deleting one column silently wrote dates onto the wrong
activities across every node in the project (audit C3). These tests reorder,
delete and shuffle columns and assert the import is still correct - or refuses.

Requires a migrated PostgreSQL database, like the API smoke tests.
"""
from __future__ import annotations

import os
from datetime import date
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

pytestmark = pytest.mark.skipif(
    os.environ.get("IPTT_TEST_DB_READY") != "1",
    reason="requires a migrated PostgreSQL database",
)

PROJECT_ID = 12


@pytest.fixture
def db():
    from app.db import SessionLocal

    session = SessionLocal()
    try:
        yield session
    finally:
        session.rollback()
        session.close()


@pytest.fixture
def workbook_bytes(db) -> bytes:
    from app.services.excel import build_workbook

    return build_workbook(db, PROJECT_ID)


def _load(content: bytes):
    return load_workbook(BytesIO(content), data_only=True)


def _save(wb) -> bytes:
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _parse(db, content: bytes, dry_run=True):
    from app.services.excel import parse_workbook

    return parse_workbook(
        db,
        PROJECT_ID,
        content,
        dry_run=dry_run,
        actor_username="tester",
        actor_user_id=None,
        actor_role="admin",
    )


# --- export -----------------------------------------------------------------


def test_export_has_the_expected_sheets(workbook_bytes):
    wb = _load(workbook_bytes)
    assert "Execution" in wb.sheetnames
    assert "Detail" in wb.sheetnames
    assert "How to use" in wb.sheetnames


def test_every_activity_column_carries_its_template_number(workbook_bytes):
    """The number is what the importer binds to, so it must always be present."""
    from app.services.excel import HEADER_PATTERN, IDENTITY_COLUMNS

    wb = _load(workbook_bytes)
    header = [c.value for c in wb["Execution"][1]]
    assert header[: len(IDENTITY_COLUMNS)] == list(IDENTITY_COLUMNS)

    activity_headers = header[len(IDENTITY_COLUMNS) :]
    assert activity_headers, "no activity columns exported"
    for title in activity_headers:
        assert HEADER_PATTERN.match(str(title)), f"unparseable header: {title!r}"


def test_export_covers_every_node(db, workbook_bytes):
    from app.models import Scope
    from sqlalchemy import func, select

    expected = db.scalar(
        select(func.count()).select_from(Scope).where(Scope.project_id == PROJECT_ID)
    )
    wb = _load(workbook_bytes)
    assert wb["Execution"].max_row - 1 == expected


def test_export_can_be_filtered_to_one_circle(db):
    from app.services.excel import build_workbook

    wb = _load(build_workbook(db, PROJECT_ID, circle="TN"))
    circles = {row[1] for row in wb["Execution"].iter_rows(min_row=2, values_only=True)}
    assert circles == {"TN"}


# --- round trip -------------------------------------------------------------


def test_unmodified_export_reimports_as_a_no_op(db, workbook_bytes):
    report = _parse(db, workbook_bytes)
    assert report.ok, report.errors
    assert report.changes == []
    assert report.nodes_matched > 0


def test_an_edit_is_detected_and_described(db, workbook_bytes):
    wb = _load(workbook_bytes)
    ws = wb["Execution"]
    header = [str(c.value) for c in ws[1]]
    # First editable Start column.
    start_col = next(i for i, h in enumerate(header, 1) if h.endswith("— Start"))
    ws.cell(row=2, column=start_col, value=date(2026, 4, 1))
    ws.cell(row=2, column=start_col + 1, value=date(2026, 4, 2))

    report = _parse(db, _save(wb))
    assert report.ok, report.errors
    assert report.changes
    fields = {c.field for c in report.changes}
    assert fields <= {"actual_start", "actual_finish"}


# --- THE C3 REGRESSION ------------------------------------------------------


def test_reordered_columns_still_import_correctly(db, workbook_bytes):
    """Reverse every activity column and confirm the import is unchanged.

    Under the legacy importer this scenario wrote each node's dates onto
    completely different activities without any error.
    """
    wb = _load(workbook_bytes)
    ws = wb["Execution"]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    identity_width = 4

    reordered = Workbook()
    target = reordered.active
    target.title = "Execution"
    for row in rows:
        head, tail = row[:identity_width], row[identity_width:]
        # Reverse the activity block wholesale, keeping each column's own header
        # attached to its own data.
        target.append(list(head) + list(reversed(tail)))

    report = _parse(db, _save(reordered))
    assert report.ok, report.errors
    assert report.changes == [], "reordering columns must not change any value"


def test_deleting_activity_columns_narrows_the_import_rather_than_shifting_it(
    db, workbook_bytes
):
    wb = _load(workbook_bytes)
    ws = wb["Execution"]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]

    trimmed = Workbook()
    target = trimmed.active
    target.title = "Execution"
    # Keep the identity columns and only the first six activity columns.
    for row in rows:
        target.append(list(row[:4]) + list(row[4:10]))

    report = _parse(db, _save(trimmed))
    assert report.ok, report.errors
    assert report.changes == []
    assert report.nodes_matched > 0


def test_removing_rows_narrows_the_import(db, workbook_bytes):
    wb = _load(workbook_bytes)
    ws = wb["Execution"]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]

    subset = Workbook()
    target = subset.active
    target.title = "Execution"
    target.append(rows[0])
    for row in rows[1:4]:
        target.append(row)

    report = _parse(db, _save(subset))
    assert report.ok, report.errors
    assert report.rows_read == 3
    assert report.nodes_matched == 3


# --- validation -------------------------------------------------------------


def test_a_mangled_activity_heading_is_rejected_not_guessed(db, workbook_bytes):
    wb = _load(workbook_bytes)
    ws = wb["Execution"]
    header = [str(c.value) for c in ws[1]]
    start_col = next(i for i, h in enumerate(header, 1) if h.endswith("— Start"))
    ws.cell(row=1, column=start_col, value="Actual Start")  # number stripped

    report = _parse(db, _save(wb))
    assert not report.ok
    assert any("could not be understood" in e for e in report.errors)


def test_a_missing_node_column_is_rejected(db, workbook_bytes):
    wb = _load(workbook_bytes)
    ws = wb["Execution"]
    ws.cell(row=1, column=1, value="Site")
    report = _parse(db, _save(wb))
    assert not report.ok
    assert any("No 'Node' column" in e for e in report.errors)


def test_a_finish_without_a_start_is_rejected(db, workbook_bytes):
    wb = _load(workbook_bytes)
    ws = wb["Execution"]
    header = [str(c.value) for c in ws[1]]
    start_col = next(i for i, h in enumerate(header, 1) if h.endswith("— Start"))
    ws.cell(row=2, column=start_col).value = None
    ws.cell(row=2, column=start_col + 1).value = date(2026, 4, 2)

    report = _parse(db, _save(wb))
    assert not report.ok
    assert any("needs a start date" in e for e in report.errors)


def test_a_finish_before_its_start_is_rejected(db, workbook_bytes):
    """The exact corruption the migration found four instances of in live data."""
    wb = _load(workbook_bytes)
    ws = wb["Execution"]
    header = [str(c.value) for c in ws[1]]
    start_col = next(i for i, h in enumerate(header, 1) if h.endswith("— Start"))
    ws.cell(row=2, column=start_col, value=date(2026, 5, 15))
    ws.cell(row=2, column=start_col + 1, value=date(2026, 3, 20))

    report = _parse(db, _save(wb))
    assert not report.ok
    assert any("precedes start" in e for e in report.errors)


def test_an_unparseable_date_is_reported_with_its_location(db, workbook_bytes):
    wb = _load(workbook_bytes)
    ws = wb["Execution"]
    header = [str(c.value) for c in ws[1]]
    start_col = next(i for i, h in enumerate(header, 1) if h.endswith("— Start"))
    ws.cell(row=2, column=start_col, value="sometime in April")

    report = _parse(db, _save(wb))
    assert not report.ok
    assert any("not a date" in e and "Row 2" in e for e in report.errors)


def test_unknown_nodes_are_warned_about_not_silently_dropped(db, workbook_bytes):
    wb = _load(workbook_bytes)
    ws = wb["Execution"]
    ws.cell(row=ws.max_row + 1, column=1, value="NOT_A_REAL_NODE")

    report = _parse(db, _save(wb))
    assert report.ok
    assert any("not in this project" in w for w in report.warnings)


def test_a_sheet_with_no_execution_tab_is_rejected(db):
    wb = Workbook()
    wb.active.title = "Sheet1"
    report = _parse(db, _save(wb))
    assert not report.ok
    assert any("no 'Execution' sheet" in e for e in report.errors)


def test_a_non_workbook_is_rejected_cleanly(db):
    report = _parse(db, b"this is not a spreadsheet")
    assert not report.ok
    assert any("could not be opened" in e for e in report.errors)


def test_plan_columns_are_ignored_even_if_edited(db, workbook_bytes):
    """Planned dates are the baseline. Editing them in the sheet must do nothing."""
    wb = _load(workbook_bytes)
    ws = wb["Execution"]
    header = [str(c.value) for c in ws[1]]
    plan_col = next(i for i, h in enumerate(header, 1) if "Plan Finish" in h)
    ws.cell(row=2, column=plan_col, value=date(2030, 1, 1))

    report = _parse(db, _save(wb))
    assert report.ok, report.errors
    assert report.changes == []


# --- application ------------------------------------------------------------


def test_dry_run_writes_nothing(db, workbook_bytes):
    from sqlalchemy import select

    from app.models import Scope, Task, TaskExecution

    wb = _load(workbook_bytes)
    ws = wb["Execution"]
    header = [str(c.value) for c in ws[1]]
    start_col = next(i for i, h in enumerate(header, 1) if h.endswith("— Start"))
    node_id = ws.cell(row=2, column=1).value
    ws.cell(row=2, column=start_col).value = date(2026, 4, 1)
    ws.cell(row=2, column=start_col + 1).value = None

    number = int(header[start_col - 1].split("·")[0].strip())
    before = db.execute(
        select(TaskExecution.actual_start)
        .join(Task, Task.id == TaskExecution.task_id)
        .join(Scope, Scope.id == TaskExecution.scope_id)
        .where(Scope.node_id == node_id, Task.template_task_number == number)
    ).scalar_one()

    report = _parse(db, _save(wb), dry_run=True)
    assert report.ok and report.changes
    db.rollback()

    after = db.execute(
        select(TaskExecution.actual_start)
        .join(Task, Task.id == TaskExecution.task_id)
        .join(Scope, Scope.id == TaskExecution.scope_id)
        .where(Scope.node_id == node_id, Task.template_task_number == number)
    ).scalar_one()
    assert after == before


def test_applying_writes_values_audit_rows_and_recomputed_delay(db, workbook_bytes):
    from sqlalchemy import func, select

    from app.models import AuditLog, Scope, Task, TaskExecution

    wb = _load(workbook_bytes)
    ws = wb["Execution"]
    header = [str(c.value) for c in ws[1]]
    start_col = next(i for i, h in enumerate(header, 1) if h.endswith("— Start"))
    number = int(header[start_col - 1].split("·")[0].strip())
    node_id = ws.cell(row=2, column=1).value

    ws.cell(row=2, column=start_col, value=date(2026, 4, 1))
    ws.cell(row=2, column=start_col + 1, value=date(2026, 4, 10))

    audits_before = db.scalar(select(func.count()).select_from(AuditLog))

    report = _parse(db, _save(wb), dry_run=False)
    assert report.ok, report.errors
    db.flush()

    execution, task, scope = db.execute(
        select(TaskExecution, Task, Scope)
        .join(Task, Task.id == TaskExecution.task_id)
        .join(Scope, Scope.id == TaskExecution.scope_id)
        .where(Scope.node_id == node_id, Task.template_task_number == number)
    ).one()

    assert execution.actual_start == date(2026, 4, 1)
    assert execution.actual_finish == date(2026, 4, 10)
    assert execution.status == "Completed"

    # Delay must be working days against task.planned_finish (decisions 1 + 2).
    from app.domain.calendar import WorkingCalendar
    from app.domain.delay import compute_delay_days

    expected = compute_delay_days(
        WorkingCalendar.from_session(db), task.planned_finish, date(2026, 4, 10), scope.circle
    )
    assert execution.delay_days == expected

    audits_after = db.scalar(select(func.count()).select_from(AuditLog))
    assert audits_after > audits_before, "the import must write audit rows (audit H12)"

    db.rollback()


def test_a_file_with_any_error_applies_nothing(db, workbook_bytes):
    """One bad cell must not leave the other rows half-written."""
    wb = _load(workbook_bytes)
    ws = wb["Execution"]
    header = [str(c.value) for c in ws[1]]
    start_col = next(i for i, h in enumerate(header, 1) if h.endswith("— Start"))

    ws.cell(row=2, column=start_col, value=date(2026, 4, 1))
    ws.cell(row=2, column=start_col + 1, value=date(2026, 4, 5))
    # Row 3 is invalid.
    ws.cell(row=3, column=start_col, value=date(2026, 5, 15))
    ws.cell(row=3, column=start_col + 1, value=date(2026, 3, 1))

    report = _parse(db, _save(wb), dry_run=False)
    assert not report.ok
    assert not db.new, "no rows should have been staged for insert"
    db.rollback()
